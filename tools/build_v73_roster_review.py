#!/usr/bin/env python3
"""Stage supplied European names for review, never for career ownership."""
import argparse
import csv
import re
import unicodedata
from collections import defaultdict
from pathlib import Path

from docx import Document


def key(value):
    value = unicodedata.normalize("NFKD", value).encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9]", "", value)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("squads", type=Path)
    ap.add_argument("clubs", type=Path)
    ap.add_argument("output", type=Path)
    args = ap.parse_args()
    with args.clubs.open(encoding="utf-8", newline="") as handle:
        clubs = list(csv.DictReader(handle, delimiter="\t"))
    by_name = {key(c["name"]): c for c in clubs}
    aliases = {
        "fcsb": "steauabucuresti", "astragiurgiu": "astraploiesti",
    }
    doc = Document(args.squads)
    entries = []
    missing = []
    for i, para in enumerate(doc.paragraphs[:-2]):
        if ";" not in doc.paragraphs[i + 1].text or not doc.paragraphs[i + 2].text.startswith("Source basis:"):
            continue
        name = para.text.strip().replace("  [ADD]", "")
        club = by_name.get(aliases.get(key(name), key(name)))
        if club is None:
            missing.append(name)
            continue
        for player in doc.paragraphs[i + 1].text.split(";"):
            player = player.strip()
            if player:
                entries.append((club["id"], player, key(player)))
    if missing or len(entries) != 2321:
        raise SystemExit(f"Unmatched clubs {missing}; parsed {len(entries)} / 2321 names")
    owners = defaultdict(set)
    for club_id, _, identity in entries:
        owners[identity].add(club_id)
    # The supplied document combines season-wide rather than opening-day rosters.
    # An apparent duplicate is a review flag, not proof of two simultaneous owners.
    flags = {identity for identity, ids in owners.items() if len(ids) > 1}
    args.output.parent.mkdir(parents=True, exist_ok=True)
    with args.output.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle, delimiter="\t", lineterminator="\n")
        writer.writerow(("club_id", "player_name", "identity_key", "review_status"))
        for club_id, player, identity in entries:
            status = "CROSS_CLUB" if identity in flags else "DRAFT"
            if identity == "derekboateng" and club_id.endswith("dnipro-dnipropetrovsk"):
                status = "DATE_CONFLICT"
            writer.writerow((club_id, player, identity, status))
    print(f"{len(entries)} entries; {len(flags)} cross-club identities; {len(missing)} unmatched clubs")


if __name__ == "__main__":
    main()
